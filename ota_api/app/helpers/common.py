import pytz
import json
import traceback
from fastapi import Request
from datetime import datetime
from dateutil.parser import parse as dataParser
from decimal import Decimal, ROUND_DOWN
from databases.database import OtaDbSession
from app.models.otadb.APILog import APILog
from app.helpers.constants import BD_TIMEZONE
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
from typing import Optional
import uuid
import os
from typing import List, Any, Tuple

try:
    from jose import jwt as _jose_jwt, JWTError as _JWTError
    from config import JWT_SECRET_KEY as _JWT_SECRET, JWT_ALGORITHM as _JWT_ALG
    _JWT_AVAILABLE = True
except Exception:
    _JWT_AVAILABLE = False
import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import Paragraph, Table, TableStyle, SimpleDocTemplate, Spacer
from reportlab.lib import colors
from app.models.otadb.ExportDownloadManager import ExportDownloadManager
from config import BASE_DIR, UPLOAD_DIR

# inporting service
from app.services.log import log_exception

def common_response(status: int = 200, message: str = "Success", data: dict = {}, errors: dict = None):
    return {
        "status": status,
        "message": message,
        "errors": errors,
        "data": data
    }

def write_log(excetion, source, type = 'error'):
    print(f'{source} || {type}: {excetion}')
    log_exception(excetion, source, type)

def format_date(date, format: str = "%Y-%m-%d"):
    if date:
        try:
            if isinstance(date, str):
                parsed_date = dataParser(date, tzinfos={'Asia/Dhaka': pytz.timezone('Asia/Dhaka')})
            elif isinstance(date, datetime):
                parsed_date = date.astimezone(BD_TIMEZONE)
            else:
                parsed_date = datetime.now(BD_TIMEZONE)
            return parsed_date.strftime(format)
        except:
            return None
    return None

def password_strength_check(password: str):
    return (
            len(password) >= 8
            and any(char.isalpha() for char in password)
            and any(char.isdigit() for char in password)
        )

def get_ota_db_session():
    db = OtaDbSession()
    try:
        yield db
    finally:
        db.close()

        
# making a dictionary into an instance
class DictionaryToInstance:
    def __init__(self, **kwargs):
        for key, value in kwargs.items():
            setattr(self, key, value)

def row_to_dict(row):
    return {column: getattr(row, column) for column in row.keys()}

def custom_float(number):
    number_str = str(number)
    if len(number_str) >= 2:
        modified_number_str = number_str[:-2] + '.' + number_str[-2:]
        modified_number_decimal = Decimal(modified_number_str).quantize(Decimal('0.00'), rounding=ROUND_DOWN)
        return str(modified_number_decimal)
    else:
        return number_str

def formatted_float(value, default=0.00):
    try:
        return round(float(value), 2)
    except (TypeError, ValueError, AttributeError):
        return default
    
def handle_float(value):
    if isinstance(value, float):
        if np.isnan(value) or np.isinf(value):
            return None
        else:
            return float(value)
    else:
        return value
    
class DecimalEncoder(json.JSONEncoder):
    def default(self, obj):
        if isinstance(obj, Decimal):
            return float(obj)
        return super(DecimalEncoder, self).default(obj)
    
def custom_round(number, decimal_places=0):
    factor = 10 ** decimal_places
    rounded = int(number * factor) / factor
    return rounded

def get_client_ip(request: Request) -> str:
    if "X-Forwarded-For" in request.headers:
        return request.headers["X-Forwarded-For"].split(",")[0]
    else:
        return request.client.host

def get_optional_user_id(request: Request) -> Optional[str]:
    """
    Silently extracts `user_id` from the Bearer JWT without doing a full
    token_validation (no Redis / DB hit).  Returns None for guests or on
    any decoding failure — never raises.
    """
    if not _JWT_AVAILABLE:
        return None
    try:
        auth_header = request.headers.get("Authorization", "")
        if not auth_header.lower().startswith("bearer "):
            return None
        token = auth_header.split(" ", 1)[1].strip()
        payload = _jose_jwt.decode(
            token, _JWT_SECRET, algorithms=[_JWT_ALG],
            options={"verify_exp": False}   # match existing token_validation behaviour
        )
        return payload.get("user_id") or payload.get("id")
    except Exception:
        return None
    
def get_error_info(exception):
    tracebacks = traceback.extract_tb(exception.__traceback__)
    filename, line_number, func_name, _ = tracebacks[-1]
    error_message = str(exception)
    return {
        "filename": filename,
        "line_number": line_number,
        "func_name": func_name,
        "error_message": error_message,
        "traceback": traceback.format_tb(exception.__traceback__)
    }

def convert_million_to_crore(million):  
    if not isinstance(million, (float, int, Decimal)):
        return million
    return formatted_float(million / 10)

def format_str_to_date(date_str):
    try:
        if date_str is not None:
            return datetime.strptime(date_str, '%d%m%Y').strftime('%Y-%m-%d')
        return date_str
    except ValueError as e:
        return None
    
def save_log(log_data: dict):
    db = next(get_ota_db_session())
    try:
        log_entry = APILog(
            timestamp=log_data.get("timestamp"),
            method=log_data.get("method"),
            url=log_data.get("url"),
            client_ip=log_data.get("client_ip"),
            headers=log_data.get("headers"),
            query_params=log_data.get("query_params"),
            request_body=log_data.get("request_body"),
            raw_request_body=log_data.get("raw_request_body", ""),
            user_agent=log_data.get("user_agent"),
            response_status=log_data.get("response_status"),
            response_size=log_data.get("response_size"),
            response_body=log_data.get("response_body"),
            process_time=log_data.get("process_time"),
            error_message=log_data.get("error_message"),
            error_details=log_data.get("error_details"),
            username=log_data.get("username")
        )
        db.add(log_entry)
        db.commit()
    except Exception as e:
        db.rollback()
        print(f"Error saving log: {e}")
    finally:
        db.close()

def create_export_record(
    db: Session,
    user_id: uuid.UUID,
    title: str,
    file_name: str,
    type_: str,
    url: Optional[str] = None,
    remarks: Optional[str] = None,
    status: str = "pending",
) -> Optional[ExportDownloadManager]:
    try:
        new_export = ExportDownloadManager(
            user_id=user_id,
            title=title,
            file_name=file_name,
            url=url,
            remarks=remarks,
            status=status,
            type=type_,            
            created_at = datetime.now(BD_TIMEZONE),
            updated_at=None
        )

        db.add(new_export)
        db.commit()
        db.refresh(new_export)

        return new_export

    except SQLAlchemyError as e:
        db.rollback()
        write_log(e, source="create_export_record", type="error")
        return None

def create_report_excel(file_path: str, title: str, data: list[list]) -> bool:
    if not data or len(data) < 2:
        return False

    headers = data[0]
    rows = data[1:]
    df = pd.DataFrame(rows, columns=headers)

    # create excel writer using openpyxl engine
    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name="Sheet1", index=False, startrow=2)
        ws = writer.sheets["Sheet1"]

        # write title across merged cells in row 1
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.font = Font(size=14, bold=True)
        title_cell.alignment = Alignment(horizontal="center")

        # add total row for "Amount" col if it exists
        amount_col = None
        for idx, col in enumerate(headers):
            if col.lower() == "amount":
                amount_col = idx
                break

        if amount_col is not None:
            total_amount = 0.0
            for val in df.iloc[:, amount_col]:
                try:
                    total_amount += float(val)
                except Exception:
                    pass

            total_row_idx = df.shape[0] + 4  # data rows + header + title + 1
            ws.cell(row=total_row_idx, column=amount_col, value="Total").font = Font(bold=True)
            ws.cell(row=total_row_idx, column=amount_col + 1, value=total_amount).font = Font(bold=True)

    return True

class PageNumCanvas(Canvas):
    def __init__(self, *args, **kwargs):
        Canvas.__init__(self, *args, **kwargs)
        self._page_number = 0
        self._saved_page_states = []

    def showPage(self):
        self._page_number += 1
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        # add page numbers to all pages
        total_pages = self._page_number
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(total_pages)
            Canvas.showPage(self)
        Canvas.save(self)

    def draw_page_number(self, total_pages):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setStrokeColorRGB(0, 0, 0)
        self.setLineWidth(0.5)

        page_width, _ = A4
        margin = 20
        footer_y = 15 * mm
        text = f"Powered By: API  |  Page {self._pageNumber} of {total_pages}"

        self.line(margin, footer_y + 8, page_width - margin, footer_y + 8)
        self.drawRightString(page_width - margin, footer_y, text)
        self.restoreState()

def create_report_pdf(file_path: str, title: str, data: List[List]) -> bool:
    styles = getSampleStyleSheet()

    # custom styles
    normal_small = ParagraphStyle(name='NormalSmall', parent=styles['Normal'], fontSize=8, leading=10)
    right_aligned_small = ParagraphStyle(name='RightSmall', parent=styles['Normal'], fontSize=8, leading=10, alignment=2)
    bold_left_aligned = ParagraphStyle(name='BoldLeft', parent=styles['Title'], fontSize=11, leading=13, alignment=0)
    title_style = ParagraphStyle(name='TitleCentered', parent=styles['Title'], fontSize=12, leading=14, alignment=1)

    if not data or len(data) < 2:
        return False

    raw_headers = data[0]

    # wrap headers
    wrapped_data = [[Paragraph(f"<b>{col}</b>", normal_small) for col in raw_headers]]

    # wrap rows
    for row in data[1:]:
        wrapped_row = [Paragraph(str(cell), normal_small) for cell in row]
        wrapped_data.append(wrapped_row)

    # add total row if "amount" exists
    amount_index = None
    for i, h in enumerate(raw_headers):
        if h.lower() == "amount":
            amount_index = i
            break

    if amount_index is not None:
        total_amount = 0.0
        for row in data[1:]:
            try:
                val = row[amount_index]
                total_amount += float(val) if val else 0.0
            except Exception:
                pass

        total_row = [""] * len(raw_headers)
        total_row[amount_index - 1 if amount_index > 0 else 0] = Paragraph("<b>Total</b>", normal_small)
        total_row[amount_index] = Paragraph(f"<b>{total_amount:.2f}</b>", normal_small)
        wrapped_data.append(total_row)

    # column widths(some has fixed widths, others will share remaining space)
    page_width, page_height = A4
    margin = 20
    usable_width = page_width - margin * 2

    fixed_widths = {
        "sl.": 30,
        "paymenttype": 65,
        "paymentmethod": 75,
        "source": 40,
    }

    dynamic_columns = [h for h in raw_headers if h.lower() not in fixed_widths]
    remaining_width = usable_width - sum(fixed_widths.get(h.lower(), 0) for h in raw_headers)
    dynamic_width = remaining_width / max(1, len(dynamic_columns))

    col_widths = [
        fixed_widths.get(h.lower(), dynamic_width)
        for h in raw_headers
    ]

    # Build document
    doc = SimpleDocTemplate(
        file_path,
        pagesize=A4,
        rightMargin=margin,
        leftMargin=margin,
        topMargin=40,
        bottomMargin=60,
    )

    elements = [
        Paragraph(f"<b>{title}</b>", title_style),
        Spacer(1, 12),
        Table(wrapped_data, colWidths=col_widths, style=TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ]),
        ),
    ]

    doc.build(elements, canvasmaker=PageNumCanvas)
    return True



def generate_user_file(record_id: int, data: List[Tuple[Any]], file_type: str, clean_title: str):
    db: Session = next(get_ota_db_session())
    record = None
    try:
        record = db.query(ExportDownloadManager).filter_by(id=record_id).first()
        if not record:
            record = db.query(ExportDownloadManager).filter_by(id=record_id).first() # Retry once
            if not record:
                write_log(f"{file_type.upper()} generation failed: Record not found", "generate_user_file_background")
                return

        record.status = "processing"
        db.commit()

        file_name = f"{clean_title}_({datetime.now(BD_TIMEZONE).strftime('%d%m%Y%H%M%S')}).{file_type}"
        record.file_name = file_name
        
        output_folder = os.path.join(BASE_DIR, UPLOAD_DIR, "users")
        os.makedirs(output_folder, exist_ok=True)

        output_path = os.path.join(output_folder, file_name)

        if not data:
            record.status = "failed"
            record.remarks = f"No data provided for {file_type.upper()} generation"
            db.commit()
            return

        # Define fields to export
        field_names = [
            "name", "username", "email", "mobile", "status", "created_at"
        ]
        headers = ["SL."] + [''.join(word.capitalize() for word in col.split('_')) for col in field_names]

        rows = []
        for idx, row_obj in enumerate(data, start=1):
            row = [idx]
            for col in field_names:
                val = getattr(row_obj, col)
                if isinstance(val, datetime):
                    val = val.strftime("%Y-%m-%d %H:%M:%S")
                row.append(val if val is not None else "")
            rows.append(row)

        table_data = [headers] + rows

        success = False
        if file_type.lower() == "pdf":
            # Reusing existing PDF generator if suitable, otherwise might need a specific one. 
            # create_report_pdf is generic enough based on previous inspection
            success = create_report_pdf(file_path=output_path, title=clean_title, data=table_data)
        elif file_type.lower() == "xlsx":
            success = create_report_excel(file_path=output_path, title=clean_title, data=table_data)
        else:
            record.status = "failed"
            record.remarks = f"Unsupported file type: {file_type}"
            db.commit()
            return

        if success:
            record.url = output_path
            record.status = "completed"
            record.remarks = None
        else:
            record.status = "failed"
            record.remarks = f"Unknown {file_type.upper()} generation error"

        db.commit()

    except Exception as e:
        if record:
            record.status = "failed"
            record.remarks = f"Exception during {file_type.upper()} processing: {str(e)}"
            db.commit()
        write_log(e, "generate_user_file_background")
    finally:
        db.close()
